from pathlib import Path

from docx import Document
from docx.shared import Pt

from resumes.models import ResumeDocument


def export_docx(document: ResumeDocument, path: Path) -> Path:
    output = Document()
    styles = output.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    output.add_heading(document.name or "候选人", level=0)
    if document.contact_line:
        output.add_paragraph(document.contact_line)
    output.add_heading(f"求职目标：{document.target_title}", level=1)
    _paragraph_section(output, "个人简介", document.summary)
    _paragraph_section(output, "专业能力", document.professional_overview)
    if document.skills:
        output.add_paragraph("技能关键词：" + "、".join(document.skills))
    _experience_section(output, "项目经历", document.projects)
    _experience_section(output, "比赛经历", document.competitions)
    _experience_section(output, "实习经历", document.internships)
    if document.education:
        output.add_heading("教育经历", level=1)
        for item in document.education:
            output.add_heading(
                " | ".join(
                    value
                    for value in (item.institution, item.degree, item.major, item.period)
                    if value
                ),
                level=2,
            )
            for bullet in item.bullets:
                output.add_paragraph(bullet.content, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    output.save(path)
    return path


def _paragraph_section(output: Document, title: str, content: str) -> None:
    if content:
        output.add_heading(title, level=1)
        output.add_paragraph(content)


def _experience_section(output: Document, title: str, items) -> None:
    if not items:
        return
    output.add_heading(title, level=1)
    for item in items:
        output.add_heading(
            " | ".join(
                value for value in (item.name, item.organization, item.role, item.period) if value
            ),
            level=2,
        )
        if item.summary:
            output.add_paragraph(item.summary)
        if item.technologies:
            output.add_paragraph("技术/能力：" + "、".join(item.technologies))
        for bullet in item.bullets:
            output.add_paragraph(bullet.content, style="List Bullet")
