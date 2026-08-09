from pathlib import Path
from profile.models import ExperienceSection, ProfileClaim, ProfileResult

from docx import Document
from docx.shared import Pt


def export_docx(result: ProfileResult, path: Path) -> Path:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.font.size = Pt(10.5)
    document.add_heading("用户五维画像", level=0)
    document.add_paragraph(f"生成时间：{result.generated_at.isoformat()}")

    document.add_heading("个人信息", level=1)
    document.add_paragraph(result.personal_introduction.overview)
    for item in result.personal_introduction.items:
        document.add_paragraph(f"{item.label}：{item.value}", style="List Bullet")

    _add_experience(document, "项目经历", result.project_experiences)
    _add_experience(document, "比赛经历", result.competition_experiences)
    _add_experience(document, "实习经历", result.internship_experiences)

    document.add_heading("学校履历", level=1)
    document.add_paragraph(result.education_history.overview)
    for entry in result.education_history.entries:
        document.add_heading(entry.institution, level=2)
        document.add_paragraph(entry.overview)
        document.add_paragraph(f"学历：{entry.degree or '资料未提供'}")
        document.add_paragraph(f"专业：{entry.major or '资料未提供'}")
        document.add_paragraph(f"时间：{entry.period or '资料未提供'}")
        document.add_paragraph(f"平均成绩：{entry.average_score or '资料未提供'}")
        document.add_paragraph(f"排名：{entry.ranking or '资料未提供'}")
        document.add_paragraph(f"综合评价：{entry.evaluation or '资料未提供'}")
        document.add_paragraph(f"语言成绩：{'、'.join(entry.language_scores) or '资料未提供'}")
        document.add_paragraph(f"核心课程：{'、'.join(entry.courses) or '资料未提供'}")
        _add_claims(document, entry.honors + entry.campus_experiences)

    document.add_heading("质量审计", level=1)
    document.add_paragraph(f"引用覆盖率：{result.audit.citation_coverage:.0%}")
    for warning in result.audit.warnings:
        document.add_paragraph(warning, style="List Bullet")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def _add_experience(document: Document, title: str, section: ExperienceSection) -> None:
    document.add_heading(title, level=1)
    document.add_paragraph(section.overview)
    for entry in section.entries:
        document.add_heading(entry.name, level=2)
        document.add_paragraph(entry.summary)
        document.add_paragraph(f"单位/组织：{entry.organization or '资料未提供'}")
        document.add_paragraph(f"时间：{entry.period or '资料未提供'}")
        document.add_paragraph(f"角色：{entry.role or '资料未提供'}")
        _add_claims(document, entry.details + entry.outcomes)


def _add_claims(document: Document, claims: list[ProfileClaim]) -> None:
    for claim in claims:
        label = "事实" if claim.basis_type.value == "fact" else "推断"
        sources = "、".join(dict.fromkeys(item.file_name for item in claim.evidence_refs)) or "无"
        document.add_paragraph(
            f"[{label} {claim.confidence:.0%}] {claim.content}（来源：{sources}）",
            style="List Bullet",
        )
