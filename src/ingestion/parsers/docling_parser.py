from pathlib import Path
from uuid import uuid4

from core.settings import Settings, get_settings
from ingestion.models import BoundingBox, ParsedDocument, SourceSpan
from ingestion.ocr import OcrService
from ingestion.parsers.base import BaseParser


class DoclingParser(BaseParser):
    def __init__(self, settings: Settings | None = None):
        self.settings = settings or get_settings()
        self.ocr = OcrService(self.settings)
        self._converter = None

    @property
    def converter(self):
        if self._converter is None:
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import PdfPipelineOptions
            from docling.document_converter import DocumentConverter, PdfFormatOption

            pipeline_options = PdfPipelineOptions()
            pipeline_options.do_ocr = self.settings.ocr_enabled
            pipeline_options.do_table_structure = True
            self._converter = DocumentConverter(
                format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)}
            )
        return self._converter

    def parse(self, path: Path, document_id: str, file_name: str) -> ParsedDocument:
        warnings: list[str] = []
        try:
            result = self.converter.convert(path)
            parsed = self._from_docling(result.document, document_id, file_name)
            if parsed.spans:
                return parsed
            warnings.append("Docling 未提取到文本，已切换本地解析器。")
        except Exception as exc:
            warnings.append(f"Docling 解析失败，已切换本地解析器：{exc}")
        fallback = self._fallback(path, document_id, file_name)
        return fallback.model_copy(update={"warnings": warnings + fallback.warnings})

    def _from_docling(self, document, document_id: str, file_name: str) -> ParsedDocument:
        spans: list[SourceSpan] = []
        page_count = 0
        for index, item in enumerate(getattr(document, "texts", []) or [], start=1):
            text = str(getattr(item, "text", "")).strip()
            if not text:
                continue
            page = None
            bbox = None
            provenance = list(getattr(item, "prov", []) or [])
            if provenance:
                source = provenance[0]
                page = getattr(source, "page_no", None)
                if page is not None:
                    page = int(page)
                    page_count = max(page_count, page)
                raw_bbox = getattr(source, "bbox", None)
                if raw_bbox is not None:
                    values = {
                        "left": getattr(raw_bbox, "l", getattr(raw_bbox, "left", 0)),
                        "top": getattr(raw_bbox, "t", getattr(raw_bbox, "top", 0)),
                        "right": getattr(raw_bbox, "r", getattr(raw_bbox, "right", 0)),
                        "bottom": getattr(raw_bbox, "b", getattr(raw_bbox, "bottom", 0)),
                    }
                    bbox = BoundingBox(**values)
            label = getattr(item, "label", None)
            spans.append(
                SourceSpan(
                    id=str(uuid4()),
                    document_id=document_id,
                    file_name=file_name,
                    text=text,
                    page=page,
                    paragraph=index,
                    section=str(label) if label else None,
                    bbox=bbox,
                )
            )
        if not spans:
            markdown = str(document.export_to_markdown()).strip()
            for index, block in enumerate(markdown.split("\n\n"), start=1):
                if block.strip():
                    spans.append(
                        SourceSpan(
                            id=str(uuid4()),
                            document_id=document_id,
                            file_name=file_name,
                            text=block.strip(),
                            paragraph=index,
                        )
                    )
        return ParsedDocument(
            document_id=document_id,
            file_name=file_name,
            spans=spans,
            page_count=page_count,
        )

    def _fallback(self, path: Path, document_id: str, file_name: str) -> ParsedDocument:
        extension = path.suffix.lower()
        if extension == ".pdf":
            return self._parse_pdf(path, document_id, file_name)
        if extension == ".docx":
            return self._parse_docx(path, document_id, file_name)
        if extension == ".pptx":
            return self._parse_pptx(path, document_id, file_name)
        raise RuntimeError(f"没有可用的解析器处理 {extension}。")

    def _parse_pdf(self, path: Path, document_id: str, file_name: str) -> ParsedDocument:
        import fitz

        spans: list[SourceSpan] = []
        with fitz.open(path) as document:
            if document.needs_pass:
                raise RuntimeError("PDF 已加密，请先解除密码保护。")
            for page_index, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                source_type = "text"
                if len(text) < 30 and self.settings.ocr_enabled:
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    recognized = self.ocr.recognize(pixmap.tobytes("png"))
                    text = "\n".join(item[0] for item in recognized)
                    source_type = "ocr"
                if text:
                    spans.append(
                        SourceSpan(
                            id=str(uuid4()),
                            document_id=document_id,
                            file_name=file_name,
                            text=text,
                            page=page_index,
                            source_type=source_type,
                        )
                    )
            return ParsedDocument(
                document_id=document_id,
                file_name=file_name,
                spans=spans,
                page_count=len(document),
            )

    def _parse_docx(self, path: Path, document_id: str, file_name: str) -> ParsedDocument:
        from docx import Document

        document = Document(path)
        spans: list[SourceSpan] = []
        paragraph = 0
        for item in document.paragraphs:
            if item.text.strip():
                paragraph += 1
                spans.append(
                    SourceSpan(
                        id=str(uuid4()),
                        document_id=document_id,
                        file_name=file_name,
                        text=item.text.strip(),
                        paragraph=paragraph,
                        section=item.style.name if item.style else None,
                    )
                )
        for table_index, table in enumerate(document.tables, start=1):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            text = "\n".join(row for row in rows if row.strip(" |"))
            if text:
                spans.append(
                    SourceSpan(
                        id=str(uuid4()),
                        document_id=document_id,
                        file_name=file_name,
                        text=text,
                        section=f"table-{table_index}",
                        source_type="table",
                    )
                )
        return ParsedDocument(document_id=document_id, file_name=file_name, spans=spans)

    def _parse_pptx(self, path: Path, document_id: str, file_name: str) -> ParsedDocument:
        from pptx import Presentation

        presentation = Presentation(path)
        spans: list[SourceSpan] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts).strip()
            if text:
                spans.append(
                    SourceSpan(
                        id=str(uuid4()),
                        document_id=document_id,
                        file_name=file_name,
                        text=text,
                        slide=slide_index,
                    )
                )
        return ParsedDocument(
            document_id=document_id,
            file_name=file_name,
            spans=spans,
            page_count=len(presentation.slides),
        )
