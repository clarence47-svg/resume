from pathlib import Path

from docx import Document
from docx.shared import Pt

from matching.models import JDMatchResult, TailoredTextSection

SECTION_TITLES = {
    "personal_introduction": "个人介绍",
    "professional_introduction": "专业介绍",
    "project_experiences": "项目经历",
    "competition_experiences": "比赛经历",
    "internship_experiences": "实习经历",
    "education_history": "学校履历",
}


def export_match_docx(result: JDMatchResult, path: Path) -> Path:
    document = Document()
    document.styles["Normal"].font.name = "Noto Sans CJK SC"
    document.styles["Normal"].font.size = Pt(10.5)
    document.add_heading(f"{result.jd_analysis.role_title} · JD 匹配文案", level=0)
    document.add_paragraph(f"总匹配度：{result.overall_score:.1f}｜版本：v{result.version}")
    document.add_heading("岗位说明", level=1)
    document.add_paragraph(result.job_research.role_summary)
    document.add_paragraph(
        f"核心能力：{'、'.join(result.job_research.core_capabilities) or '未识别'}"
    )
    document.add_paragraph(
        f"常见职责：{'；'.join(result.job_research.typical_responsibilities) or '未识别'}"
    )
    document.add_paragraph(f"常见工具：{'、'.join(result.job_research.common_tools) or '未识别'}")
    if result.job_research.sources:
        document.add_heading("岗位研究来源", level=2)
        for source in result.job_research.sources:
            document.add_paragraph(f"{source.title}：{source.url}", style="List Bullet")
    document.add_heading("匹配分析", level=1)
    for item in result.strengths:
        document.add_paragraph(f"优势：{item}", style="List Bullet")
    for item in result.gaps:
        document.add_paragraph(f"缺口：{item}", style="List Bullet")
    for field, title in SECTION_TITLES.items():
        section = getattr(result, field)
        document.add_heading(title, level=1)
        document.add_paragraph(section.overview.content)
        if isinstance(section, TailoredTextSection):
            for unit in section.bullets:
                document.add_paragraph(unit.content, style="List Bullet")
        else:
            for entry in section.entries:
                name = getattr(entry, "name", None) or getattr(entry, "institution", "")
                document.add_heading(name, level=2)
                document.add_paragraph(entry.tailored_summary.content)
                if hasattr(entry, "organization"):
                    document.add_paragraph(
                        "｜".join(
                            [
                                entry.organization or "资料未提供",
                                entry.period or "资料未提供",
                                entry.role or "资料未提供",
                            ]
                        )
                    )
                    document.add_paragraph(
                        f"核心技术：{'、'.join(entry.technologies) or '资料未提供'}"
                    )
                else:
                    document.add_paragraph(
                        "｜".join(
                            [
                                entry.degree or "资料未提供",
                                entry.major or "资料未提供",
                                entry.period or "资料未提供",
                            ]
                        )
                    )
                for unit in entry.bullets:
                    document.add_paragraph(unit.content, style="List Bullet")
    document.add_heading("质量审计", level=1)
    document.add_paragraph(f"证据覆盖率：{result.audit.evidence_coverage:.0%}")
    for warning in result.audit.warnings:
        document.add_paragraph(warning, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path
